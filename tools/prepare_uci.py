"""Convert the public Online Retail II workbook to the workbench CSV contract."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--country", default="United Kingdom")
    parser.add_argument("--top-skus", type=int, default=12)
    return parser


def main() -> int:
    args = build_parser().parse_args()
    if not args.input.is_file():
        raise SystemExit(f"Workbook was not found: {args.input}")
    if not 1 <= args.top_skus <= 100:
        raise SystemExit("--top-skus must be between 1 and 100")

    daily: dict[tuple[str, str, date], list[float]] = defaultdict(lambda: [0.0, 0.0])
    sku_totals: dict[str, float] = defaultdict(float)
    skipped = 0
    scanned = 0
    workbook = load_workbook(args.input, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        raw_headers = next(rows)
        headers = {str(value).strip(): index for index, value in enumerate(raw_headers)}
        invoice_column = _column(headers, "Invoice", "InvoiceNo")
        sku_column = _column(headers, "StockCode")
        quantity_column = _column(headers, "Quantity")
        date_column = _column(headers, "InvoiceDate")
        price_column = _column(headers, "Price", "UnitPrice")
        country_column = _column(headers, "Country")

        for row in rows:
            scanned += 1
            invoice = str(row[invoice_column] or "").strip()
            sku = str(row[sku_column] or "").strip()
            country = str(row[country_column] or "").strip()
            try:
                quantity = float(row[quantity_column])
                price = float(row[price_column])
                observed_on = _as_date(row[date_column])
            except (TypeError, ValueError):
                skipped += 1
                continue
            if (
                not sku
                or invoice.upper().startswith("C")
                or quantity <= 0
                or price < 0
                or country != args.country
            ):
                skipped += 1
                continue
            key = (sku, country, observed_on)
            daily[key][0] += quantity
            daily[key][1] += quantity * price
            sku_totals[sku] += quantity

    active_days: dict[str, int] = defaultdict(int)
    for sku, _country, _observed_on in daily:
        active_days[sku] += 1
    selected = {
        sku
        for sku in sorted(
            sku_totals,
            key=lambda item: (active_days[item], sku_totals[item]),
            reverse=True,
        )[: args.top_skus]
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["date", "sku_id", "location_id", "units", "price"])
        for (sku, country, observed_on), (units, revenue) in sorted(daily.items()):
            if sku not in selected:
                continue
            writer.writerow(
                [observed_on.isoformat(), sku, country, f"{units:g}", f"{revenue / units:.4f}"]
            )

    report: dict[str, Any] = {
        "source_sha256": hashlib.sha256(args.input.read_bytes()).hexdigest(),
        "output_sha256": hashlib.sha256(args.output.read_bytes()).hexdigest(),
        "scanned_rows": scanned,
        "skipped_rows": skipped,
        "selected_skus": len(selected),
        "selection_rule": "highest active-day coverage, then unit volume",
        "output_rows": sum(1 for _ in args.output.open(encoding="utf-8")) - 1,
        "country": args.country,
    }
    report_path = args.output.with_suffix(".report.json")
    report_path.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report))
    return 0


def _column(headers: dict[str, int], *candidates: str) -> int:
    for candidate in candidates:
        if candidate in headers:
            return headers[candidate]
    raise ValueError(f"Workbook is missing a required column: {' or '.join(candidates)}")


def _as_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.fromisoformat(str(value)).date()


if __name__ == "__main__":
    raise SystemExit(main())
